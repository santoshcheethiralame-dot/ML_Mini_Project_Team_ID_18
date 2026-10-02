"""Build the review slide deck (with speaker notes) from results/metrics.json.

    python scripts/make_slides.py --authors "Name1, Name2" --problem-id 7
"""
import argparse
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

DARK, GREEN, LIGHT, INK, MUTED = (RGBColor.from_string(c) for c in ("0B3D2E", "05A862", "F2F8F5", "1F2A26", "5E6B65"))
WHITE = RGBColor(255, 255, 255)
PAPER_ACC = 0.83
FONT = "Calibri"


def pct(x, d=1):
    return f"{100 * x:.{d}f}%"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, dark=False, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = DARK if dark else WHITE
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    @staticmethod
    def text(s, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = line
            r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
            r.font.color.rgb = color
            if i:
                p.space_before = Pt(size * 0.45)
        return tb

    def title(self, s, t, sub=None, dark=False):
        self.text(s, 0.7, 0.45, 11.9, 0.8, t, 32, True, WHITE if dark else DARK)
        if sub:
            self.text(s, 0.7, 1.2, 11.9, 0.5, sub, 16, False, MUTED if not dark else RGBColor(190, 220, 205))

    @staticmethod
    def card(s, x, y, w, h, fill=LIGHT):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background()
        sh.adjustments[0] = 0.06; sh.shadow.inherit = False
        return sh

    def stat(self, s, x, y, w, big, small):
        self.card(s, x, y, w, 1.35)
        self.text(s, x + 0.2, y + 0.18, w - 0.4, 0.65, big, 30, True, GREEN)
        self.text(s, x + 0.2, y + 0.85, w - 0.4, 0.4, small, 13, False, MUTED)

    def bullets(self, s, x, y, w, h, items, size=18):
        self.text(s, x, y, w, h, ["\u2022  " + i for i in items], size)

    @staticmethod
    def image(s, path, x, y, w=None, h=None):
        if os.path.exists(path):
            kw = {}
            if w: kw["width"] = Inches(w)
            if h: kw["height"] = Inches(h)
            s.shapes.add_picture(path, Inches(x), Inches(y), **kw)

    def table(self, s, x, y, w, rows, col_w=None, size=14):
        shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.42 * len(rows)))
        tb = shape.table
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = tb.cell(i, j); c.text = str(val)
                p = c.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                f = p.runs[0].font; f.size, f.name = Pt(size), FONT
                f.bold = (i == 0) or (i == 1)
                f.color.rgb = WHITE if i == 0 else INK
                c.fill.solid(); c.fill.fore_color.rgb = DARK if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
        if col_w:
            for j, cw in enumerate(col_w):
                tb.columns[j].width = Inches(cw)
        return tb


def build(M, authors, problem_id, fig_dir, team_id=""):
    st, meta = M["meta"]["stats"], M["meta"]
    t1, t2, abl = M["table_sample1"], M["table_sample2"], M["ablation"]
    best = t2[0]
    a1, a2 = abl["S1_meta"]["acc_mean"], abl["S2_nb"]["acc_mean"]
    extras = {k: abl[k]["acc_mean"] - a2 for k in ("S3_nb_sent", "S4_nb_lda", "S5_nb_lsa")}
    cm = M["confusion"]
    D = Deck()
    fig = lambda n: os.path.join(fig_dir, n)

    # 1 title
    s = D.slide(dark=True, notes="Introduce the problem in one sentence: can we predict from launch-time info alone whether a Kickstarter campaign will hit its goal?")
    D.text(s, 0.8, 2.2, 11.5, 1.6, "Predicting Kickstarter\nProject Success", 54, True, WHITE)
    D.text(s, 0.8, 4.4, 11.5, 0.5, "Launch-time features, leakage-safe encodings, gradient boosting", 20, False, RGBColor(190, 220, 205))
    tid = f"  |  Team #{team_id}" if team_id else ""
    pid = f"  |  Problem #{problem_id}" if problem_id else ""
    D.text(s, 0.8, 5.6, 11.5, 0.8, [f"UE24CS352A Machine Learning Mini-Project{tid}{pid}", authors], 16, False, WHITE)
    if meta["data_source"] != "real":
        D.text(s, 0.8, 0.5, 11.5, 0.5, "DRAFT - BUILT FROM SYNTHETIC SMOKE-TEST DATA, NOT REAL RESULTS", 16, True, RGBColor(255, 120, 120))

    # 2 problem
    s = D.slide(notes="Binary classification. Key constraint: only information known at launch. That makes the model useful BEFORE a creator spends effort, and avoids leakage from backers/pledged amount.")
    D.title(s, "The problem", "Replicating and extending Khosla, Reinecke & Wittenbrink (Stanford)")
    D.bullets(s, 0.7, 2.1, 7.0, 4.5, [
        "Input: what a creator knows at launch (goal, category, country, currency, location, dates, name, one-line blurb)",
        "Output: will the campaign reach its funding goal? (binary)",
        "Why it matters: creators can tune goal and framing; the platform can prioritise promising projects",
        f"Paper's best result: about {pct(PAPER_ACC, 0)} accuracy",
    ], 20)
    D.card(s, 8.3, 2.1, 4.3, 2.9)
    D.text(s, 8.6, 2.4, 3.8, 0.5, "Excluded on purpose", 18, True, DARK)
    D.text(s, 8.6, 3.0, 3.8, 2.5, ["backers count", "amount pledged so far", "staff pick / spotlight", "anything after launch"], 17, False, INK)

    # 3 dataset
    s = D.slide(notes=f"Source: Web Robots monthly scrapes. They are cumulative, so we keep the latest row per project id, keep only successful/failed, and dedupe on (name, blurb, launch, deadline) like the paper. Majority-class baseline is {pct(st['baseline_acc'])}, so accuracy must be read against that.")
    D.title(s, "Dataset", "Web Robots Kickstarter scrape, finished campaigns only")
    D.stat(s, 0.7, 1.9, 2.9, f"{st['n_final']:,}", "projects after cleaning")
    D.stat(s, 3.8, 1.9, 2.9, pct(st["success_rate"]), "succeeded")
    D.stat(s, 6.9, 1.9, 2.9, f"{st['n_train']:,} / {st['n_test']:,}", f"train / test ({st['split']} 70-30)")
    D.stat(s, 10.0, 1.9, 2.6, pct(st["baseline_acc"]), "majority-class accuracy")
    D.image(s, fig("eda.png"), 1.6, 3.55, h=3.6)

    # 4 features
    s = D.slide(notes="Three families. The target encodings and NB probability use labels, so they are out-of-fold. Goal means per level do not use labels so they are fitted on all training rows.")
    D.title(s, "Feature engineering", "Most of the performance comes from the first two groups")
    cols = [("Metadata", ["log goal (USD)", "campaign length", "launch month / weekday", "name & blurb length", "one-hot parent category & category"]),
            ("Out-of-fold target encodings", ["currency, country, category, parent, location, location type", "mean goal per level", "goal relative to level mean"]),
            ("Text (blurb)", ["TF-IDF \u2192 Naive Bayes probability", "VADER sentiment", "20-topic LDA", "50-d LSA embedding"])]
    for i, (h, items) in enumerate(cols):
        x = 0.7 + i * 4.1
        D.card(s, x, 1.9, 3.9, 3.9)
        D.text(s, x + 0.25, 2.1, 3.4, 0.6, h, 20, True, DARK)
        D.bullets(s, x + 0.25, 2.9, 3.4, 3.6, items, 16)

    # 5 leakage
    s = D.slide(dark=True, notes="This is the point to emphasise in Q&A. If a row's encoding is computed with its own label, the model can memorise it and accuracy is inflated. Out-of-fold: split train into 5 folds; encode each fold using the other 4. Test rows use encoders fitted on the whole train set. There is a unit test for this in tests/.")
    D.title(s, "Avoiding leakage", dark=True)
    D.text(s, 0.8, 1.7, 11.7, 4.8, [
        "Target encodings and the Naive-Bayes probability depend on the label.",
        "On the training set they are computed out-of-fold (5 folds): each row is encoded using only the other folds.",
        "Test and demo rows use encoders fitted on the full training set.",
        "Duplicates removed before splitting, so no project appears in both train and test.",
        "Unit tests check the encoding guard and that no post-launch columns are used.",
    ], 22, False, WHITE)

    # 6 models
    s = D.slide(notes="Same model families as the paper. Linear probability models (OLS/ridge/lasso) regress y and threshold at 0.5. Differences from the paper: VADER instead of CoreNLP (no Java server), LSA instead of 3.6GB pretrained Word2Vec, sklearn ReLU MLP instead of SELU Keras net.")
    D.title(s, "Models and protocol")
    D.bullets(s, 0.7, 1.7, 6.2, 5, [
        "Gradient boosting (LightGBM)",
        "Random forest",
        "Neural net (1 hidden layer, 25 units)",
        "Linear SVM, logistic regression",
        "Ridge / OLS / lasso linear probability models",
    ], 20)
    D.card(s, 7.4, 1.7, 5.3, 3.3)
    D.text(s, 7.7, 1.95, 4.8, 0.5, "Protocol", 20, True, DARK)
    D.bullets(s, 7.7, 2.6, 4.8, 3.6, ["Balanced sample weights", "Fit on train only; test touched once", "Accuracy, F1, per-class P/R, ROC-AUC",
                                      f"Text ablation over {len(abl['S1_meta']['accs'])} seeds"], 16)

    # 7 results table
    s = D.slide(notes=f"Compare to the majority baseline ({pct(st['baseline_acc'])}) and to the paper ({pct(PAPER_ACC, 0)}). Class 1 is 'successful'. Check whether precision on failed projects (P0) is the weak spot, as in the paper.")
    D.title(s, "Results: metadata + Naive Bayes feature", f"Test set, {st['n_test']:,} projects")
    rows = [["Model", "Acc", "F1", "P1", "P0", "R1", "R0", "AUC"]] + [[r["model"], *(f"{r[k]:.3f}" for k in ("A", "F1", "P1", "P0", "R1", "R0", "AUC"))] for r in t2]
    D.table(s, 0.7, 1.9, 8.2, rows, [2.0] + [0.9] * 7, 15)
    D.card(s, 9.4, 1.9, 3.3, 3.4)
    D.text(s, 9.65, 2.15, 2.8, 0.6, f"{best['model']}", 22, True, DARK)
    D.text(s, 9.65, 2.8, 2.8, 2.4, [f"{pct(best['A'])} accuracy", f"F1 {best['F1']:.3f}", f"{100 * (best['A'] - st['baseline_acc']):+.1f} pts vs baseline"], 18)

    # 8 model comparison
    s = D.slide(notes="Adding the NB feature lifts every model. Top models are close, so features matter more than the model.")
    D.title(s, "Features matter more than the model", f"NB feature: {pct(a1)} \u2192 {pct(a2)} accuracy ({100 * (a2 - a1):+.1f} pts, LightGBM)")
    D.image(s, fig("models.png"), 1.7, 1.9, w=9.8)

    # 9 ablation
    s = D.slide(notes=(f"Be precise, not vague. Beyond metadata + NB: sentiment {100 * extras['S3_nb_sent']:+.2f} points, "
                       f"LDA {100 * extras['S4_nb_lda']:+.2f}, LSA {100 * extras['S5_nb_lsa']:+.2f}. Sentiment adds nothing "
                       "measurable on a one-line blurb; the topic and embedding blocks each add about half a point - so do "
                       "not say the text features add nothing. The defensible claim: the blurb supports one well-regularised "
                       "text feature (the NB probability) and the extra blocks refine it slightly. Boxplots are across "
                       "random seeds of the boosting model."))
    D.title(s, "Do extra text features help?", "Gain over Metadata + NB (accuracy points): " + ", ".join(f"{abl[k]['label'].split('+ ')[-1]} {100 * v:+.2f}" for k, v in extras.items()))
    D.image(s, fig("ablation.png"), 2.0, 2.0, w=9.3)

    # 10 error analysis
    s = D.slide(notes="Confusion matrix is for gradient boosting. nb_prob, target encodings and goal-relative features dominate the gain importance.")
    D.title(s, "Error analysis and feature importance")
    D.image(s, fig("confusion.png"), 0.7, 1.7, w=5.0)
    D.image(s, fig("importance.png"), 6.1, 1.7, w=6.6)
    D.text(s, 0.7, 6.3, 11.9, 0.8, f"{cm[1][0]:,} successes predicted as failures; {cm[0][1]:,} failures predicted as successes.", 16, False, MUTED)

    # 11 demo
    s = D.slide(notes="Live demo: streamlit run app.py. Change the goal and watch the probability; open the contributions chart to show which features drive it.")
    D.title(s, "Live demo", "streamlit run app.py")
    D.bullets(s, 0.7, 2.1, 11.5, 4.5, [
        "Enter name, blurb, category, country, goal and dates",
        "Get a success probability from the trained gradient-boosting model",
        "What-if curve: how the prediction changes as the goal changes",
        "Per-prediction feature contributions (TreeSHAP) for explainability",
    ], 22)

    # 12 conclusions
    s = D.slide(dark=True, notes="Close with limitations: only launch-time metadata and a one-line blurb; no story, images, creator history. A temporal split is a stricter test. Next step: transformer sentence embeddings.")
    D.title(s, "Conclusions", dark=True)
    D.text(s, 0.8, 1.7, 11.7, 5, [
        f"{best['model']} is best: {pct(best['A'])} accuracy, F1 {best['F1']:.3f} (paper: about {pct(PAPER_ACC, 0)}).",
        f"Out-of-fold Naive-Bayes blurb probability is the main text feature ({100 * (a2 - a1):+.2f} pts); "
        f"sentiment adds {100 * extras['S3_nb_sent']:+.2f} pts, LDA {100 * extras['S4_nb_lda']:+.2f} and LSA {100 * extras['S5_nb_lsa']:+.2f}.",
        "Leakage-safe encodings and de-duplication keep the estimate honest.",
        "Future work: temporal validation, transformer embeddings, project story and image features.",
    ], 22, False, WHITE)

    return D.prs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--metrics", default="results/metrics.json")
    ap.add_argument("--out", default="results/slides.pptx")
    ap.add_argument("--authors", default="Team members")
    ap.add_argument("--problem-id", default="")
    ap.add_argument("--team-id", default="")
    a = ap.parse_args()
    M = json.load(open(a.metrics))
    prs = build(M, a.authors, a.problem_id, os.path.join(os.path.dirname(a.metrics), "figures"), a.team_id)
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    prs.save(a.out)
    print(f"wrote {a.out} ({len(prs.slides)} slides)")
    if M["meta"]["data_source"] != "real":
        print("WARNING: metrics come from SYNTHETIC data; the title slide is marked as a draft")


if __name__ == "__main__":
    main()
